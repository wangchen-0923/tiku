"""删除 4.3 后的题库健康度重算 + 编码规则影响检查。只读。"""

from __future__ import annotations

import re
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

Q = Path(r"D:\tiku\docs\题目资料.xlsx")
OUT = Path(r"D:\tiku\docs\题库-删4.3后健康度.txt")

CN = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6, "七": 7}

lines: list[str] = []


def p(s: str = "") -> None:
    lines.append(s)


main_q: dict[str, dict[str, int]] = defaultdict(lambda: defaultdict(int))
wb = load_workbook(Q, data_only=False)
for sheet, key in (("选择题", "选择"), ("判断题", "判断"), ("Sheet2", "填空")):
    ws = wb[sheet]
    for r in list(ws.iter_rows(values_only=True))[1:]:
        if r and r[0] is not None:
            ch = str(r[1]).strip() if len(r) > 1 and r[1] is not None else "?"
            main_q[ch][key] += 1
wb.close()

# 任务定义（第X章：任务Y 标题）
TASKS = {
    "1.1": "分析轴的力学性能", "1.2": "识别非合金钢", "1.3": "选择合金钢",
    "1.4": "判别铸铁", "1.5": "识别非铁金属", "1.6": "轴的热处理",
    "1.7": "选择非金属材料",
    "2.1": "平面连杆机构", "2.2": "凸轮机构", "2.3": "间歇机构",
    "3.1": "销、螺纹、键连接", "3.2": "联轴器与离合器",
    "4.1": "轴承", "4.2": "轴及轴上零件",
    "5.1": "带传动", "5.2": "链传动", "5.3": "螺旋传动",
    "5.4": "齿轮传动", "5.5": "蜗杆传动", "5.6": "齿轮系",
}
# 知识点数（按清单，去重后每任务下的知识点个数）
KP_COUNT = {
    "1.1": 2, "1.2": 2, "1.3": 1, "1.4": 1, "1.5": 2, "1.6": 2, "1.7": 2,
    "2.1": 3, "2.2": 2, "2.3": 2,
    "3.1": 3, "3.2": 2,
    "4.1": 2, "4.2": 2,
    "5.1": 3, "5.2": 2, "5.3": 3, "5.4": 2, "5.5": 2, "5.6": 1,
}

DELETE = {"4.3"}

p("=" * 90)
p("删除 4.3（69 题）后的题库健康度")
p("=" * 90)
p(f"{'章节':<7}{'任务名称':<20}{'选择':>6}{'判断':>6}{'填空':>6}{'合计':>7}"
  f"{'知识点':>7}  {'每知识点题量':>12}  状态")
p("-" * 90)

total_all = 0
total_after = 0
for ch in sorted(main_q, key=lambda s: [int(x) for x in s.split(".")]):
    row = main_q[ch]
    tot = sum(row.values())
    total_all += tot
    if ch in DELETE:
        p(f"{ch:<7}{'(已删除)':<20}{row['选择']:>6}{row['判断']:>6}{row['填空']:>6}"
          f"{tot:>7}      -            -   ❌ 删除")
        continue
    total_after += tot
    n_kp = KP_COUNT.get(ch, 0)
    per = f"{tot/n_kp:.0f}" if n_kp else "-"
    if tot < 30:
        status = "⚠ 偏少"
    elif row["填空"] == 0:
        status = "⚠ 无填空题"
    else:
        status = "OK"
    p(f"{ch:<7}{TASKS.get(ch,'?'):<20}{row['选择']:>6}{row['判断']:>6}{row['填空']:>6}"
      f"{tot:>7}{n_kp:>7}{per:>12}  {status}")

p("-" * 90)
p(f"{'合计（删前）':<27}{'':>18}{total_all:>7}")
p(f"{'合计（删后）':<27}{'':>18}{total_after:>7}")
p(f"删除题量: {total_all - total_after}")

# 章维度
p()
p("=" * 90)
p("章维度题量（删除 4.3 后）")
p("=" * 90)
by_chapter: dict[int, int] = defaultdict(int)
for ch, row in main_q.items():
    if ch in DELETE:
        continue
    by_chapter[int(ch.split(".")[0])] += sum(row.values())
for c in sorted(by_chapter):
    p(f"    第{c}章  {by_chapter[c]:>5} 题")
p(f"    合计    {sum(by_chapter.values()):>5} 题")

# 第4章特别说明
p()
p("★ 第4章（轴承与轴）删除 4.3 后的情况：")
c4 = sum(v for k, v in by_chapter.items() if k == 4)
p(f"    剩余题量: {c4} 题（4.1 轴承 + 4.2 轴及轴上零件）")
p(f"    任务数: 2（原本清单也只有 2 个任务，删除后与清单一致）")

# 编码规则影响
p()
p("=" * 90)
p("编码规则影响检查（我的 DDL 限制 vs 你的编码形式）")
p("=" * 90)
samples = ["MJJC-C1-T1-K1", "MJJC.C1.T1.K1", "1.1.1", "K-1-1-1", "机械基础-C1-T1-K1"]
pat_old = re.compile(r"^[A-Za-z0-9._-]{1,80}$")
for s in samples:
    p(f"    {s:<26} 原规则(字母数字._-) {'通过' if pat_old.match(s) else '拒绝'}")
p()
p("说明：DDL 里 code 字段是 varchar(80)，本身不限字符；")
p("      限制来自我写的正则校验。中文编码也应允许 → 放宽为：")
p("      长度 1-80，只禁止空白、斜杠与前导点，其余放开。")

OUT.write_text("\n".join(lines), encoding="utf-8")
print(f"OK -> {OUT} ({len(lines)} lines)")
