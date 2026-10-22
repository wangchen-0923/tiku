"""回归：知识库文档文本抽取层（需求 3 / `P1-18b`）。

被测对象：`server/apps/knowledge/extract.py`（`.docx` / `.pdf` / `.txt` → `ExtractedDoc`）

★ 本脚本的铁律：
  1. **测试数据现场生成，不落库**——全部 fixtures 写在 `tempfile.TemporaryDirectory()`
     里，`finally` 清干净，并断言临时目录真的消失（不留文件）；
  2. **每条断言都打印实测值**（字符数 / 页数 / 命中的编码 / warnings 原文），
     否则「通过」二字没有任何信息量；
  3. 每条断言都能回答「它在什么情况下会红」。

★ 关于 PDF fixture 的字体选择（踩过的坑，写下来免得后人再踩）：

  需求里建议用 reportlab 内置 CID 字体 `STSong-Light` 造中文 PDF。
  实测**不可用于本回归**：`UnicodeCIDFont("STSong-Light")` 写出的 PDF
  没有 `ToUnicode` CMap，PyMuPDF 抽出来是一串替换字符（`'����е����'`），
  「中文文本层」这条断言会**假红**（红的原因与抽取层无关）。
  所以这里改用 PyMuPDF 自带的 `Font("cjk")`（Droid Sans Fallback，
  随 wheel 分发、不依赖系统字体）+ `TextWriter` 写文本层，
  抽出来是干净的中文；取不到时退化成 Windows 系统字体 `simhei.ttf`。
"""

from __future__ import annotations

import os
import sys
import tempfile
import traceback

# ══════════════════════════════════════════════════════════════
#  环境（★ DB 端口必须是 15433：settings 默认的 5432 是别的项目的库）
# ══════════════════════════════════════════════════════════════
sys.path.insert(0, r"D:\tiku\.devdeps")
sys.path.insert(0, r"D:\tiku\server")
os.environ.pop("ZF_DB_ENGINE", None)
os.environ["DJANGO_SETTINGS_MODULE"] = "zhifou.settings.local"
os.environ["ZF_CACHE"] = "locmem"
os.environ["ZF_DB_HOST"] = "127.0.0.1"
os.environ["ZF_DB_PORT"] = "15433"
os.environ["ZF_DB_NAME"] = "zhifou"
os.environ["ZF_DB_USER"] = "zhifou"
os.environ["ZF_DB_PASSWORD"] = "zhifou"

try:  # GBK 控制台：不 reconfigure 就会在中文/警告原文处乱码
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:  # pragma: no cover
    pass

import django  # noqa: E402

django.setup()

from apps.common.exceptions import ErrorCode, ValidationFailed, ZfError  # noqa: E402
from apps.knowledge.extract import (  # noqa: E402
    DEFAULT_MAX_BYTES,
    SCAN_PDF_WARNING,
    SOURCE_DOCX,
    SOURCE_PDF,
    SOURCE_TXT,
    extract_document,
)

PASS: list[str] = []
FAIL: list[str] = []


def check(name: str, ok: bool, detail: str = "") -> None:
    (PASS if ok else FAIL).append(name)
    mark = "[OK]  " if ok else "[FAIL]"
    print(f"  {mark} {name}" + (f"\n         实测：{detail}" if detail else ""))


def run(name: str, fn) -> None:
    """跑一条用例；用例内部抛异常 → 记 FAIL（而不是让整个脚本崩掉）。"""
    try:
        fn()
    except Exception as exc:  # noqa: BLE001
        check(name, False, f"未预期异常 {type(exc).__name__}: {exc}")
        traceback.print_exc()


# ══════════════════════════════════════════════════════════════
#  fixture 生成
# ══════════════════════════════════════════════════════════════

CN_TITLE = "第一章 机械基础"
CN_HEAD2 = "1.1 齿轮传动的失效形式"
CN_BODY = "齿轮传动是最常见的机械传动形式，其失效形式包括齿面点蚀、胶合、磨损与轮齿折断。"
CN_BODY2 = "点蚀多发生在闭式软齿面齿轮的节线附近。"
CN_TXT = "机械基础知识点：齿轮传动\n失效形式：点蚀、胶合、磨损、折断\n检查方法：目视、磁粉探伤"

#: PDF 用**短**句（见 `make_pdf_cjk` 的行宽守卫）。
CN_PDF_P1 = "齿轮传动的失效形式：点蚀、胶合、磨损"
CN_PDF_P2 = "点蚀多发生在节线附近"


def _write(path: str, data: bytes) -> str:
    with open(path, "wb") as fh:
        fh.write(data)
    return path


def _read(path: str) -> bytes:
    with open(path, "rb") as fh:
        return fh.read()


def make_docx(path: str) -> str:
    """含「标题 1 / 标题 2 / 正文 / 表格」的中文 docx。"""
    from docx import Document

    doc = Document()
    doc.add_heading(CN_TITLE, level=1)          # 样式名 Heading 1
    doc.add_paragraph(CN_BODY)
    doc.add_paragraph(CN_HEAD2, style="Heading 2")  # 显式样式名
    doc.add_paragraph(CN_BODY2)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "失效形式"
    table.cell(0, 1).text = "典型部位"
    table.cell(1, 0).text = "点蚀"
    table.cell(1, 1).text = "节线附近"
    doc.add_paragraph("")  # 空段落：应被忽略且不影响 text 与 blocks 的拼接不变式
    doc.save(path)
    return path


def make_docx_image_only(path: str) -> str:
    """只有一张图片、没有任何文字段落的 docx（模拟「截图版」知识文档）。"""
    from docx import Document
    from docx.shared import Mm

    doc = Document()
    png = _png_bytes()
    tmp_png = os.path.join(os.path.dirname(path), "_tmp_img.png")
    _write(tmp_png, png)
    doc.add_picture(tmp_png, width=Mm(40))
    doc.save(path)
    return path


def make_pdf_cjk(path: str, pages: list[str]) -> str:
    """含中文**文本层**的 PDF（每页一段中文，单行）。

    ★★ 行宽守卫（实测踩到的坑）：`TextWriter` 不自动换行，
      写超出页面右边界（A4 宽 595pt、左边距 72pt）的字**会被 MuPDF 裁掉**，
      抽取时就少最后几个字——回归会出现「看起来是抽取层少字」的假红。
      所以这里硬性限制每行字符数：14pt × 30 字 = 420pt，留足余量。
      超了直接抛，逼调用方改短句，而不是让断言以难查的方式变红。
    """
    import pymupdf

    font = _cjk_font()
    if font is None:
        raise RuntimeError("找不到可用的中文字体（pymupdf Font('cjk') 与系统字体都失败）")
    doc = pymupdf.open()
    for text in pages:
        if len(text) > _PDF_MAX_LINE_CHARS:
            doc.close()
            raise ValueError(
                f"PDF fixture 单行 {len(text)} 字超过 {_PDF_MAX_LINE_CHARS} 字，"
                "超出页宽会被裁掉：请把测试句改短"
            )
        page = doc.new_page()
        writer = pymupdf.TextWriter(page.rect)
        writer.append((72, 700), text, font=font, fontsize=14)
        writer.write_text(page)
    data = doc.tobytes()
    doc.close()
    return _write(path, data)


def make_pdf_scanned(path: str) -> str:
    """扫描件模拟：只放一张图片，**不放任何文本层**。"""
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_image(pymupdf.Rect(60, 60, 360, 260), stream=_png_bytes())
    data = doc.tobytes()
    doc.close()
    return _write(path, data)


def _png_bytes() -> bytes:
    import pymupdf

    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 80, 50))
    pix.clear_with(190)
    return pix.tobytes("png")


#: PDF 每行最多写多少字（14pt × 30 字 = 420pt，A4 可用宽度约 523pt）。
_PDF_MAX_LINE_CHARS = 30


def _cjk_font():
    import pymupdf

    try:
        return pymupdf.Font("cjk")  # 随包分发的 Droid Sans Fallback
    except Exception:  # noqa: BLE001
        for candidate in (
            r"C:\Windows\Fonts\simhei.ttf",
            r"C:\Windows\Fonts\STSONG.TTF",
            r"C:\Windows\Fonts\msyh.ttc",
        ):
            if os.path.exists(candidate):
                try:
                    return pymupdf.Font(fontfile=candidate)
                except Exception:  # noqa: BLE001
                    continue
    return None


def expect_reject(fn) -> tuple[bool, str]:
    """调用 fn，期望抛 `ZfError`；返回 (是否按预期拒绝, 实测描述)。

    ★ 把 `extra` 拍平成 `k=v` 串再打印：字典 repr 里是 `'max_bytes': 1024`，
      断言「max_bytes=1024」会永远为假——那种断言看着在测，实际测不到。
    """
    try:
        result = fn()
    except ZfError as exc:
        extra = " ".join(f"{k}={v}" for k, v in (exc.extra or {}).items())
        return True, (
            f"已拒绝：type={type(exc).__name__} code={exc.zf_code} "
            f"status={exc.status_code} message={exc.detail!r} extra=[{extra}]"
        )
    except Exception as exc:  # noqa: BLE001
        return False, f"抛了**非** ZfError：{type(exc).__name__}: {exc}"
    return False, f"没有拒绝，返回了 source_type={result.source_type} char_count={result.char_count}"


# ══════════════════════════════════════════════════════════════
#  用例
# ══════════════════════════════════════════════════════════════


def main() -> int:  # noqa: C901 - 回归脚本，平铺更好读
    tmp = tempfile.TemporaryDirectory(prefix="zf_kb_extract_")
    root = tmp.name
    tmpdir_gone = False
    try:
        print("=" * 78)
        print("知识库文本抽取回归（P1-18b）")
        print(f"临时目录：{root}")
        print("=" * 78)

        # ── 1. .docx：标题层级 ──
        print("\n[1] .docx 标题层级（P1-18c 要按层级切条，层级丢了就补不回来）")
        docx_path = make_docx(os.path.join(root, "知识库样例.docx"))
        docx_bytes = _read(docx_path)

        def case_docx():
            r = extract_document("知识库样例.docx", docx_bytes)
            levels = [b.level for b in r.blocks]
            kinds = [b.kind for b in r.blocks]
            print(f"         text={r.text!r}")
            print(f"         warnings={r.warnings}")
            check(
                "1.1 docx 识别为 DOCX 且非空",
                r.source_type == SOURCE_DOCX and not r.is_empty,
                f"source_type={r.source_type} char_count={r.char_count} "
                f"non_whitespace={r.non_whitespace_count} byte_size={r.byte_size}",
            )
            check(
                "1.2 docx 标题带 # / ## 前缀，层级分别为 1 和 2",
                "# " + CN_TITLE in r.text and "## " + CN_HEAD2 in r.text,
                f"levels={levels} kinds={kinds} heading_count={r.heading_count}",
            )
            check(
                "1.3 docx 段落数（非空块数）== 6（2 标题 + 2 正文 + 2 表格行）",
                r.paragraph_count == 6,
                f"paragraph_count={r.paragraph_count} blocks={len(r.blocks)}",
            )
            check(
                "1.4 docx 表格内容也被抽到（表格行 kind=table_row）",
                "失效形式 | 典型部位" in r.text and "点蚀 | 节线附近" in r.text
                and kinds.count("table_row") == 2,
                f"表格行={[b.text for b in r.blocks if b.kind == 'table_row']}",
            )
            check(
                "1.5 不变式：'\\n'.join(blocks) == text（含空段落也不能破坏）",
                "\n".join(b.text for b in r.blocks) == r.text,
                f"join_len={len(chr(10).join(b.text for b in r.blocks))} text_len={r.char_count}",
            )
            check(
                "1.6 docx 的 pdf 专属字段为空（page_count / encoding 均 None）",
                r.page_count is None and r.encoding is None and r.page_char_counts == [],
                f"page_count={r.page_count} encoding={r.encoding} page_char_counts={r.page_char_counts}",
            )
            check(
                "1.7 正常 docx **不**产生 warning",
                r.warnings == [],
                f"warnings={r.warnings}",
            )

        run("docx 标题层级", case_docx)

        # ── 2. .docx：全图片 ──
        print("\n[2] .docx 全是图片（有内容说没内容 = 最坏的错，必须有 warning）")
        img_docx_path = make_docx_image_only(os.path.join(root, "截图版.docx"))
        img_docx_bytes = _read(img_docx_path)

        def case_docx_image():
            r = extract_document("截图版.docx", img_docx_bytes)
            print(f"         warnings={r.warnings}")
            check(
                "2.1 docx 全图片：text 为空且 image_count >= 1",
                r.is_empty and r.image_count >= 1,
                f"char_count={r.char_count} is_empty={r.is_empty} image_count={r.image_count}",
            )
            check(
                "2.2 docx 全图片：warnings 明确说「未取到文本 / 图片 / OCR」",
                any(("未取到文本" in w and "图片" in w and "OCR" in w) for w in r.warnings),
                f"warnings={r.warnings}",
            )

        run("docx 全图片", case_docx_image)

        # ── 3. .pdf：中文文本层 ──
        print("\n[3] .pdf 含中文文本层（逐页抽取）")
        pdf_path = make_pdf_cjk(
            os.path.join(root, "知识库样例.pdf"), [CN_PDF_P1, CN_PDF_P2]
        )
        pdf_bytes = _read(pdf_path)

        def case_pdf():
            r = extract_document("知识库样例.pdf", pdf_bytes)
            print(f"         text={r.text!r}")
            print(f"         warnings={r.warnings} page_char_counts={r.page_char_counts}")
            check(
                "3.1 pdf 识别为 PDF，逐页取到 2 页中文（逐字相同，不丢尾字）",
                r.source_type == SOURCE_PDF
                and r.page_count == 2
                and CN_PDF_P1 in r.text
                and CN_PDF_P2 in r.text,
                f"source_type={r.source_type} page_count={r.page_count} "
                f"char_count={r.char_count} page_char_counts={r.page_char_counts} "
                f"第1页命中={CN_PDF_P1 in r.text} 第2页命中={CN_PDF_P2 in r.text}",
            )
            check(
                "3.2 pdf blocks 带 1 基页码（AI 引文要标页码）",
                [b.page for b in r.blocks] == [1, 2]
                and all(b.kind == "page" for b in r.blocks),
                f"pages={[b.page for b in r.blocks]} kinds={[b.kind for b in r.blocks]}",
            )
            check(
                "3.3 不变式：'\\n'.join(blocks) == text",
                "\n".join(b.text for b in r.blocks) == r.text,
                f"blocks={len(r.blocks)} text_len={r.char_count}",
            )
            check(
                "3.4 有文本层的 pdf 不报扫描件，且 docx/txt 专属字段为空",
                not any("扫描件" in w for w in r.warnings)
                and r.encoding is None
                and r.paragraph_count is None,
                f"warnings={r.warnings} encoding={r.encoding} paragraph_count={r.paragraph_count}",
            )

        run("pdf 中文文本层", case_pdf)

        # ── 4. .pdf：扫描件（无文本层）★ 最重要的一条 ──
        print("\n[4] .pdf 扫描件模拟（只有图片、无文本层）——诚实边界")
        scan_path = make_pdf_scanned(os.path.join(root, "扫描件.pdf"))
        scan_bytes = _read(scan_path)

        def case_pdf_scanned():
            r = extract_document("扫描件.pdf", scan_bytes)
            print(f"         text={r.text!r} non_whitespace={r.non_whitespace_count}")
            print(f"         warnings 原文={r.warnings}")
            check(
                "4.1 扫描件：text 空 / 近空（strip 后为空）",
                r.is_empty and r.non_whitespace_count == 0,
                f"char_count={r.char_count} non_whitespace={r.non_whitespace_count} "
                f"page_count={r.page_count} page_char_counts={r.page_char_counts}",
            )
            check(
                "4.2 扫描件：warnings 含「扫描件」+ 明确「需要 OCR」+ 指到 P3-09",
                any(
                    "扫描件" in w and "OCR" in w and "P3-09" in w
                    for w in r.warnings
                ),
                f"warnings={r.warnings}",
            )
            check(
                "4.3 扫描件：warning 与模块常量 SCAN_PDF_WARNING 逐字一致（前端按字面分支）",
                SCAN_PDF_WARNING in r.warnings,
                f"SCAN_PDF_WARNING={SCAN_PDF_WARNING!r} in warnings={SCAN_PDF_WARNING in r.warnings}",
            )
            check(
                "4.4 扫描件不算「失败」：正常返回 ExtractedDoc（不抛异常），由上层提示老师",
                r.source_type == SOURCE_PDF and r.page_count == 1,
                f"source_type={r.source_type} page_count={r.page_count} warnings={len(r.warnings)}",
            )

        run("pdf 扫描件", case_pdf_scanned)

        # ── 5. .txt：三种编码必须解出同一段文本 ★ 编码探测核心判据 ──
        print("\n[5] .txt 中文编码探测（utf-8 / utf-8-sig / gb18030 必须同文）")
        txt_variants = {
            "utf-8": CN_TXT.encode("utf-8"),
            "utf-8-sig": CN_TXT.encode("utf-8-sig"),
            "gb18030": CN_TXT.encode("gb18030"),
        }
        results = {}

        def case_txt_encodings():
            for label, data in txt_variants.items():
                path = _write(os.path.join(root, f"样例_{label}.txt"), data)
                r = extract_document(f"样例_{label}.txt", _read(path))
                results[label] = r
                print(
                    f"         写出编码={label:<10} 字节数={len(data):<5} "
                    f"命中编码={r.encoding!r} char_count={r.char_count} "
                    f"paragraph_count={r.paragraph_count} warnings={r.warnings}"
                )
                print(f"           取到文本={r.text!r}")

            texts = {label: r.text for label, r in results.items()}
            check(
                "5.1 三种编码解出的文本**完全一致**且等于原文",
                len(set(texts.values())) == 1 and set(texts.values()) == {CN_TXT},
                f"各编码文本={texts}",
            )
            check(
                "5.2 命中编码正确区分：无 BOM utf-8 → 'utf-8'；带 BOM → 'utf-8-sig'；GBK → 'gb18030'",
                results["utf-8"].encoding == "utf-8"
                and results["utf-8-sig"].encoding == "utf-8-sig"
                and results["gb18030"].encoding == "gb18030",
                f"utf-8→{results['utf-8'].encoding!r} "
                f"utf-8-sig→{results['utf-8-sig'].encoding!r} "
                f"gb18030→{results['gb18030'].encoding!r}",
            )
            check(
                "5.3 ★ 顺序反了的证伪：这 3 个文件的文本里都不能出现 gb18030 误解出的乱码字符",
                all("鏈" not in r.text and "锟" not in r.text for r in results.values()),
                f"各文本={ {k: v.text for k, v in results.items()} }",
            )
            check(
                "5.4 txt blocks == text 的非空行（空行不进 blocks）",
                [b.text for b in results["utf-8"].blocks]
                == [ln for ln in results["utf-8"].text.split("\n") if ln.strip()],
                f"blocks={[b.text for b in results['utf-8'].blocks]} "
                f"lines={[ln for ln in results['utf-8'].text.split(chr(10)) if ln.strip()]}",
            )
            check(
                "5.5 txt 的页数为 None（页数是 pdf 专属）",
                results["utf-8"].page_count is None
                and results["utf-8"].page_count is None,
                f"page_count={results['utf-8'].page_count} "
                f"paragraph_count={results['utf-8'].paragraph_count}",
            )

        run("txt 编码探测", case_txt_encodings)

        # ── 6. .txt：换行归一 + 全空白 ──
        print("\n[6] .txt 换行归一与全空白文件")

        def case_txt_crlf():
            raw = "行一\r\n行二\r行三\r\n".encode("utf-8")
            r = extract_document("换行.txt", raw)
            print(f"         text={r.text!r} paragraph_count={r.paragraph_count}")
            check(
                "6.1 CRLF / CR 全部归一为 LF（否则不同机器切行结果不一致）",
                r.text == "行一\n行二\n行三\n" and r.paragraph_count == 3,
                f"text={r.text!r} paragraph_count={r.paragraph_count} char_count={r.char_count}",
            )

        run("txt 换行归一", case_txt_crlf)

        def case_txt_blank():
            r = extract_document("空白.txt", "   \n\t\n \r\n".encode("utf-8"))
            print(f"         warnings={r.warnings} is_empty={r.is_empty}")
            check(
                "6.2 全空白 txt：不抛异常，但必须给「没有可导入内容」warning",
                r.is_empty and any("没有可导入的" in w for w in r.warnings),
                f"char_count={r.char_count} is_empty={r.is_empty} warnings={r.warnings}",
            )

        run("txt 全空白", case_txt_blank)

        # ── 7. .txt：编码回退分支 ──
        print("\n[7] .txt 编码全失败 → 兜底 + 「可能有乱码」warning")

        def case_txt_fallback():
            bad = b"\xff\xff\xfe\x00\x81\x80"  # 三种严格编码都解不开
            r = extract_document("乱码.txt", bad)
            print(f"         命中编码={r.encoding!r} text={r.text!r}")
            print(f"         warnings 原文={r.warnings}")
            check(
                "7.1 回退命中 utf-8/replace，且 warning 同时含「编码回退」与「乱码」",
                r.encoding == "utf-8/replace"
                and any("编码回退" in w and "乱码" in w for w in r.warnings),
                f"encoding={r.encoding!r} warnings={r.warnings}",
            )

        run("txt 编码回退", case_txt_fallback)

        # ── 8. 反例：扩展名与内容不符 ──
        print("\n[8] 反例：格式判定（只看扩展名会判错的那几类）")
        text_bytes = CN_TXT.encode("utf-8")
        pdf_head = pdf_bytes
        docx_head = docx_bytes

        def case_mismatch_docx_ext():
            ok, detail = expect_reject(
                lambda: extract_document("假文档.docx", pdf_head)
            )
            check(
                "8.1 .docx 后缀 + PDF 内容 → 拒绝（PK 魔数不符）",
                ok and "不符" in str(detail),
                detail,
            )

        run(".docx 后缀 / PDF 内容", case_mismatch_docx_ext)

        def case_mismatch_pdf_ext():
            ok, detail = expect_reject(
                lambda: extract_document("假文档.pdf", text_bytes)
            )
            check(
                "8.2 .pdf 后缀 + 纯文本内容 → 拒绝（无 %PDF- 魔数）",
                ok and "不符" in str(detail),
                detail,
            )

        run(".pdf 后缀 / 纯文本内容", case_mismatch_pdf_ext)

        def case_mismatch_txt_ext():
            ok, detail = expect_reject(
                lambda: extract_document("假文本.txt", pdf_head)
            )
            check(
                "8.3 .txt 后缀 + PDF 内容 → 拒绝（无魔数的 txt 也不许收二进制）",
                ok and "不符" in str(detail),
                detail,
            )

        run(".txt 后缀 / PDF 内容", case_mismatch_txt_ext)

        def case_mismatch_pdf_ext_docx():
            ok, detail = expect_reject(
                lambda: extract_document("假表格.pdf", docx_head)
            )
            check(
                "8.4 .pdf 后缀 + docx 内容 → 拒绝",
                ok and "不符" in str(detail),
                detail,
            )

        run(".pdf 后缀 / docx 内容", case_mismatch_pdf_ext_docx)

        def case_broken_docx():
            broken = b"PK\x03\x04" + b"\x00" * 200  # 有魔数，不是合法 zip
            try:
                extract_document("坏文档.docx", broken)
                check("8.5 PK 魔数但 zip 损坏 → INGEST.PARSE_FAILED", False, "没有抛出异常")
            except ZfError as exc:
                check(
                    "8.5 PK 魔数但 zip 损坏 → INGEST.PARSE_FAILED（422，而不是 400 参数错）",
                    exc.zf_code == ErrorCode.INGEST_PARSE_FAILED
                    and exc.status_code == 422,
                    f"code={exc.zf_code} status={exc.status_code} message={exc.detail!r}",
                )

        run("损坏 docx", case_broken_docx)

        def case_encrypted_pdf():
            """加密 PDF：**必须**拒绝，而不是抽出空文本让老师以为扫描件。

            ★ 这条分支不测就是死代码：`needs_pass` 判断只有真加密文件才走到。
            """
            import pymupdf

            doc = pymupdf.open()
            page = doc.new_page()
            page.insert_text((72, 700), "secret", fontsize=14)
            data = doc.tobytes(
                encryption=pymupdf.PDF_ENCRYPT_AES_256,
                owner_pw="owner-pw",
                user_pw="user-pw",
            )
            doc.close()
            ok, detail = expect_reject(
                lambda: extract_document("加密.pdf", data)
            )
            check(
                "8.6 加密 PDF → INGEST.PARSE_FAILED（422）且提示「已加密」",
                ok and "加密" in str(detail) and "422" in str(detail),
                detail,
            )

        run("加密 pdf", case_encrypted_pdf)

        # ── 9. 反例：空文件 / 超限 / 不支持的后缀 ──
        print("\n[9] 反例：空文件 / 超限 / 不支持的后缀")

        def case_empty():
            for filename in ("空.txt", "空.docx", "空.pdf"):
                ok, detail = expect_reject(lambda f=filename: extract_document(f, b""))
                if not (ok and "空" in str(detail)):
                    check(f"9.1 空文件 {filename} → 拒绝", False, detail)
                    return
            check(
                "9.1 空文件（.txt/.docx/.pdf 三种后缀）→ 拒绝且提示「空」",
                True,
                "三个空文件都被 ValidationFailed 拒绝，message 均含「空」",
            )

        run("空文件", case_empty)

        # 恰好等于上限：341 个「中」（3 字节）+ 1 个 ASCII = 1024 字节
        exact = ("中" * 341 + "a").encode("utf-8")
        over = exact + b"b"

        def case_size_limit():
            assert len(exact) == 1024 and len(over) == 1025, (len(exact), len(over))
            allowed = extract_document("刚好.txt", exact, max_bytes=1024)
            check(
                "9.2 恰好等于上限（1024B）→ 通过（上限含等号）",
                allowed.char_count == 342,
                f"byte_size={allowed.byte_size} max_bytes=1024 char_count={allowed.char_count}",
            )
            ok, detail = expect_reject(
                lambda: extract_document("超限.txt", over, max_bytes=1024)
            )
            check(
                "9.3 超上限 1 字节（1025B > 1024B）→ 拒绝且 extra 带上限与实测大小",
                ok and "max_bytes=1024" in str(detail) and "byte_size=1025" in str(detail),
                detail,
            )

        run("大小上限", case_size_limit)

        def case_default_limit():
            from apps.media.storage import KIND_MAX_BYTES

            check(
                "9.4 默认上限 20MB，且与 media.storage 的 KNOWLEDGE_ATTACH 一致",
                DEFAULT_MAX_BYTES == 20 * 1024 * 1024
                and DEFAULT_MAX_BYTES == KIND_MAX_BYTES["KNOWLEDGE_ATTACH"],
                f"DEFAULT_MAX_BYTES={DEFAULT_MAX_BYTES} "
                f"KIND_MAX_BYTES['KNOWLEDGE_ATTACH']={KIND_MAX_BYTES['KNOWLEDGE_ATTACH']}",
            )
            # 用**默认上限**跑一份真实 docx：证明「没传 max_bytes 也能正常抽取」
            r = extract_document("知识库样例.docx", docx_bytes)
            check(
                "9.5 不传 max_bytes（走默认 20MB）时正常抽取",
                not r.is_empty and r.char_count > 0,
                f"char_count={r.char_count} 默认上限={DEFAULT_MAX_BYTES}",
            )

        run("默认上限", case_default_limit)

        def case_bad_extension():
            ok1, d1 = expect_reject(lambda: extract_document("说明.md", text_bytes))
            ok2, d2 = expect_reject(lambda: extract_document("老卷子.doc", text_bytes))
            ok3, d3 = expect_reject(lambda: extract_document("没有后缀", text_bytes))
            check(
                "9.6 不支持的后缀 .md / .doc / 无后缀 → 拒绝（.doc 提示「另存为 .docx」）",
                ok1 and ok2 and ok3 and "另存为" in str(d2),
                f"md={d1} | doc={d2} | 无后缀={d3}",
            )

        run("不支持的后缀", case_bad_extension)

        # ── 10. 抽取层的纯度（不碰 ORM）──
        print("\n[10] 抽取层不碰数据库（纯抽取层：能离线跑、能单测）")

        def case_no_db():
            from django.test.utils import CaptureQueriesContext
            from django.db import connection

            with CaptureQueriesContext(connection) as ctx:
                extract_document("知识库样例.docx", docx_bytes)
                extract_document("知识库样例.pdf", pdf_bytes)
                extract_document("样例_utf-8.txt", txt_variants["utf-8"])
            check(
                "10.1 docx + pdf + txt 三次抽取共发起 0 条 SQL",
                len(ctx.captured_queries) == 0,
                f"SQL 条数={len(ctx.captured_queries)}",
            )

        run("不碰数据库", case_no_db)

    finally:
        tmp.cleanup()
        tmpdir_gone = not os.path.exists(root)
        # ★ 不留文件不是「顺手」——它是本回归可重复跑的前提：
        #   上一轮残留的 fixture 会让下一轮的「文件数/大小」类断言出现漂移。
        check(
            "11.1 临时目录已删除（回归不留任何文件）",
            tmpdir_gone,
            f"TemporaryDirectory={root} 清理后仍存在={not tmpdir_gone}",
        )
        print(f"\n临时目录已清理：{root} 存在={not tmpdir_gone}")

    print("\n" + "=" * 78)
    print(f"通过 {len(PASS)} 项，失败 {len(FAIL)} 项")
    if FAIL:
        print("失败清单：")
        for name in FAIL:
            print(f"  - {name}")
    print("=" * 78)
    return 1 if FAIL else 0


if __name__ == "__main__":
    sys.exit(main())
