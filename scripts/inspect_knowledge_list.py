"""检查知识点清单.xlsx 的结构 —— 输出 UTF-8 报告供开发使用。只读。"""

from __future__ import annotations

import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

SRC = Path(r"D:\tiku\docs\知识点清单.xlsx")
OUT = Path(r"D:\tiku\docs\知识点清单-结构报告.txt")

lines: list[str] = []


def p(s: str = "") -> None:
    lines.append(s)


def main() -> int:
    wb = load_workbook(SRC, data_only=False)
    p("=" * 78)
    p("知识点清单.xlsx 结构报告")
    p("=" * 78)
    p(f"工作表: {wb.sheetnames}")
    for ws in wb.worksheets:
        p(f"  [{ws.title}] state={ws.sheet_state} dims={ws.dimensions} "
          f"rows={ws.max_row} cols={ws.max_column} images={len(getattr(ws,'_images',[]))}")

    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        p()
        p("#" * 78)
        p(f"# 工作表【{ws.title}】  {len(rows)} 行 × {ws.max_column} 列")
        p("#" * 78)
        p(f"合并单元格: {[str(r) for r in ws.merged_cells.ranges][:20]}")

        p()
        p("--- 全部内容（逐行，带行号与列号）---")
        for i, row in enumerate(rows, start=1):
            cells = []
            for j, v in enumerate(row, start=1):
                if v is not None and str(v).strip():
                    cells.append(f"[{chr(64+j)}{i}] {str(v).strip()}")
            if cells:
                p("  " + "  |  ".join(cells))
            else:
                p(f"  (空行 {i})")

        # 非空计数
        p()
        p("--- 每列非空计数 ---")
        for j in range(ws.max_column):
            cnt = sum(1 for r in rows if j < len(r) and r[j] is not None and str(r[j]).strip())
            p(f"  列{j+1}: {cnt}")

    wb.close()
    OUT.write_text("\n".join(lines), encoding="utf-8")
    print(f"OK -> {OUT} ({len(lines)} lines)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
