"""题目资料试跑分析（不落库）。

目的：在真正导入前，把以下事实全部摆出来
  1. 三张表的真实题量与题型
  2. 选项解析成功率（选择题）
  3. 判断题 A/B 映射
  4. 填空分隔符与答案数量一致性
  5. 所属章节的分布（能否映射到知识树）
  6. 重复题、异常清单

产出：docs/教材/题目试跑报告.md
用法：python scripts/qa_dry_run.py
"""

from __future__ import annotations

import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:  # pragma: no cover
    pass

SRC = Path(r"D:\tiku\docs\题目资料.xlsx")
OUT = Path(r"D:\tiku\docs\教材\题目试跑报告.md")

# 选项标记：A. / A、 / A， / A： / A) / A） / A空格
OPT_MARK = re.compile(r"([A-D])\s*(?:[.．、,，:：)）]|\s)\s*")
# 填空占位符（下划线，全角半角混合）
BLANK = re.compile(r"[_＿]{2,}")
# 填空题答案分隔符：只有中文分号与逗号（★ 斜杠不分割，如 GB/T 119.6）
BLANK_SEP = re.compile(r"[；，]")


def s(v) -> str:
    return str(v).strip() if v is not None else ""


def parse_options(stem: str) -> tuple[str, list[str]]:
    """从题干中切出选项。返回 (纯题干, [选项文本...])。

    策略：找到 A→B→C→D 递增出现的标记位置，按位置切分。
    """
    marks = [(m.start(), m.group(1)) for m in OPT_MARK.finditer(stem)]
    # 只保留 A→B→C→D 严格递增的序列
    seq: list[tuple[int, str]] = []
    want = "ABCD"
    for pos, lab in marks:
        if len(seq) < 4 and lab == want[len(seq)]:
            seq.append((pos, lab))
    if len(seq) < 2:
        return stem.strip(), []
    body = stem[:seq[0][0]].strip()
    opts = []
    for i, (pos, lab) in enumerate(seq):
        end = seq[i + 1][0] if i + 1 < len(seq) else len(stem)
        txt = stem[pos:end]
        txt = re.sub(r"^[A-D]\s*(?:[.．、,，:：)）]|\s)\s*", "", txt).strip()
        opts.append(txt.strip(" ;；,，"))
    return body, opts


def main() -> int:
    wb = load_workbook(SRC, data_only=True)
    md: list[str] = []
    out: list[str] = []

    def p(x: str = "") -> None:
        out.append(x)

    p("# 题目资料 · 试跑报告（未落库）\n")
    p(f"数据源：`{SRC.name}`\n")

    # ══ 1. 选择题 ══
    ws = wb["选择题"]
    rows = list(ws.iter_rows(values_only=True))[1:]
    choice = [(s(r[0]), s(r[1]), s(r[2])) for r in rows if s(r[0])]
    p("## 一、选择题\n")
    p(f"题量：**{len(choice)}**\n")
    ok, noopt, badopt = [], [], []
    opt_counts: Counter[int] = Counter()
    for stem, ch, ans in choice:
        body, opts = parse_options(stem)
        if not opts:
            noopt.append((stem, ch, ans))
            continue
        opt_counts[len(opts)] += 1
        if len(opts) < 4:
            badopt.append((stem, ch, ans, len(opts)))
        else:
            ok.append((body, opts, ch, ans))
    p("| 项 | 数量 |")
    p("| --- | --- |")
    p(f"| 解析出 4 个选项 | {len(ok)} |")
    p(f"| 解析出 2-3 个选项 | {len(badopt)} |")
    p(f"| **完全未解析出选项** | **{len(noopt)}** |")
    p(f"| 解析成功率 | {len(ok)/len(choice)*100:.2f}% |")
    p()
    p(f"选项个数分布：{dict(sorted(opt_counts.items()))}")
    p()
    ans_c = Counter(a for _, _, a in choice)
    p(f"答案分布：{dict(sorted(ans_c.items()))}")
    p()

    # ══ 2. 判断题 ══
    ws = wb["判断题"]
    rows = list(ws.iter_rows(values_only=True))[1:]
    tf = [(s(r[0]), s(r[1]), s(r[2])) for r in rows if s(r[0])]
    p("## 二、判断题\n")
    p(f"题量：**{len(tf)}**\n")
    tf_c = Counter(a for _, _, a in tf)
    p(f"答案分布：{dict(sorted(tf_c.items()))}")
    p()
    p("> **A/B 映射风险（D-42）**：必须确认 A=正确、B=错误。"
      "下文用抽样题目交叉验证。")
    p()

    # ══ 3. 填空题 ══
    ws = wb["Sheet2"]
    rows = list(ws.iter_rows(values_only=True))[1:]
    fb = [(s(r[0]), s(r[1]), s(r[2])) for r in rows if s(r[0])]
    p("## 三、填空题\n")
    p(f"题量：**{len(fb)}**\n")
    blank_mismatch, no_blank = [], []
    for stem, ch, ans in fb:
        n_blank = len(BLANK.findall(stem))
        parts = [x for x in BLANK_SEP.split(ans) if x.strip()]
        if n_blank == 0:
            no_blank.append((stem, ch, ans))
        elif len(parts) != n_blank:
            blank_mismatch.append((stem, ch, ans, n_blank, parts))
    p("| 项 | 数量 |")
    p("| --- | --- |")
    p(f"| 占位符与答案数量一致 | {len(fb)-len(blank_mismatch)-len(no_blank)} |")
    p(f"| **占位符与答案数量不一致** | **{len(blank_mismatch)}** |")
    p(f"| 无下划线占位符 | {len(no_blank)} |")
    p()

    # ══ 4. 章节分布 ══
    p("## 四、所属章节分布\n")
    all_ch = Counter()
    for _, ch, _ in choice + tf + fb:
        all_ch[ch] += 1
    p(f"章节标记共 {len(all_ch)} 种\n")
    p("| 章节标记 | 题量 | 推测归属 |")
    p("| --- | --- | --- |")
    for ch in sorted(all_ch, key=lambda x: [int(y) for y in x.split(".") if y.isdigit()] or [0]):
        cnt = all_ch[ch]
        parts = ch.split(".")
        guess = f"第{parts[0]}章 任务{parts[1]}" if len(parts) == 2 else "待判"
        p(f"| `{ch}` | {cnt} | {guess} |")
    p()

    # ══ 5. 重复题 ══
    p("## 五、重复题检查\n")
    for label, data in (("选择题", [x[0] for x in choice]),
                        ("判断题", [x[0] for x in tf]),
                        ("填空题", [x[0] for x in fb])):
        c = Counter(data)
        dups = {k: v for k, v in c.items() if v > 1}
        p(f"- **{label}**：重复题干 {len(dups)} 组")
        for k, v in list(dups.items())[:5]:
            p(f"    - ×{v}　{k[:50]}…")
    p()

    # ══ 6. 异常清单 ══
    p("## 六、需人工处理的异常清单\n")
    p("### 6.1 选择题 · 未解析出选项\n")
    if noopt:
        p("| 章节 | 答案 | 题干 |")
        p("| --- | --- | --- |")
        for stem, ch, ans in noopt[:20]:
            p(f"| `{ch}` | {ans} | {stem[:60]} |")
    else:
        p("无。")
    p()
    p("### 6.2 填空题 · 占位符与答案数量不一致\n")
    if blank_mismatch:
        p("| 章节 | 占位符 | 答案份数 | 题干 | 答案 |")
        p("| --- | --- | --- | --- | --- |")
        for stem, ch, ans, nb, parts in blank_mismatch:
            p(f"| `{ch}` | {nb} | {len(parts)} | {stem[:44]} | {ans[:30]} |")
    else:
        p("无。")
    p()
    p("### 6.3 填空题 · 无占位符\n")
    if no_blank:
        p("| 章节 | 题干 | 答案 |")
        p("| --- | --- | --- |")
        for stem, ch, ans in no_blank[:20]:
            p(f"| `{ch}` | {stem[:50]} | {ans[:30]} |")
    else:
        p("无。")
    p()

    # ══ 7. 判断题 A/B 抽样验证 ══
    p("## 七、判断题 A/B 映射抽样验证\n")
    p("抽查前 15 题，人工核对「A 是否为正确陈述」：\n")
    p("| 答案 | 题干 |")
    p("| --- | --- |")
    for stem, ch, ans in tf[:15]:
        p(f"| {ans} | {stem[:58]} |")
    p()

    text = "\n".join(out)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(text, encoding="utf-8")

    print("=" * 74)
    print("题目资料 · 试跑分析")
    print("=" * 74)
    print(f"选择题 {len(choice)}   判断题 {len(tf)}   填空题 {len(fb)}"
          f"   合计 {len(choice)+len(tf)+len(fb)}")
    print(f"选择题选项解析成功 {len(ok)} / 失败 {len(noopt)} / 不全 {len(badopt)}")
    print(f"填空题占位符不匹配 {len(blank_mismatch)}，无占位符 {len(no_blank)}")
    print(f"章节标记 {len(all_ch)} 种")
    print(f"\nOK -> {OUT}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
