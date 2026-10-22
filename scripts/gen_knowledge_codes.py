"""从知识点清单生成课程/章/任务/知识点编码清单。

课程编码：JXJC（用户确认，方案 B）
编码格式：JXJC-C01-T01-K01（补零 + 固定前缀，可被程序解析）

★ 已确认的数据事实（从 Excel 实测，勿凭印象改）：
  - 表头在第 1 行，列位：C(索引2)=章节  D(索引3)=知识点  E(索引4)=核心内容
  - 每行同时含「上级名称 + 知识点」：如一行里既有「任务一：…」又有「1. …」
    因此解析时上级分支【不能 continue】，要让它继续走知识点解析
  - 「附录：润滑与密封」算正文 -> 编为第六章，归第三篇，章名去「附录」字样
    该章无任务层级，D 列直接是知识点，故每个知识点自成一个「任务」

产出：
  docs/教材/知识点编码清单.md    给教研过目
  docs/教材/知识点编码清单.csv   便于导入/校对（带 BOM）
  docs/教材/知识点树.json        迁移脚本的输入

用法：python scripts/gen_knowledge_codes.py
"""

from __future__ import annotations

import csv
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:  # pragma: no cover
    pass

SRC = Path(r"D:\tiku\docs\知识点清单.xlsx")
OUT_DIR = Path(r"D:\tiku\docs\教材")

COURSE_CODE = "JXJC"
COURSE_NAME = "机械基础"
GRADE_BAND = "SECONDARY_VOCATIONAL"

CN_NUM = "一二三四五六七八九十"


def cn2int(s: str) -> int:
    """中文数字转阿拉伯数字（支持一 ~ 二十）。"""
    if s == "十":
        return 10
    if s.startswith("十"):
        return 10 + CN_NUM.index(s[1]) + 1
    if "十" in s:
        a, _, b = s.partition("十")
        return (CN_NUM.index(a) + 1) * 10 + (CN_NUM.index(b) + 1 if b else 0)
    return CN_NUM.index(s) + 1


def s(v) -> str:
    return str(v).strip() if v is not None else ""


def parse() -> tuple[list[dict], list[dict]]:
    """解析 Excel，返回 (parts, chapters)。"""
    rows = list(load_workbook(SRC, data_only=True)["Sheet1"].iter_rows(values_only=True))

    parts: list[dict] = []
    chapters: list[dict] = []
    cur_part = cur_ch = cur_task = None

    for r in rows[1:]:
        C, D, E = s(r[2]), s(r[3]), s(r[4])

        # ── 篇 ──
        if re.match(r"^第.+篇[:：]", C):
            cur_part = {"raw": C, "name": C.split("：", 1)[1].strip(),
                        "chapters": []}
            parts.append(cur_part)
            cur_ch = cur_task = None
            continue

        # ── 章（含原「附录」，已按用户确认并入章序列）──
        if C.startswith("附录"):
            ch_name = C.split("：", 1)[1].strip() if "：" in C else "附录"
            marker = True
        else:
            m = re.match(r"^(第[一二三四五六七八九十]+章)[:：](.+)$", C)
            ch_name = m.group(2).strip() if m else None
            marker = bool(m)

        if marker and ch_name and not any(c["name"] == ch_name for c in chapters):
            cur_ch = {"seq": len(chapters) + 1, "name": ch_name,
                      "code_suffix": None, "tasks": [], "part": cur_part,
                      "no_task_layer": C.startswith("附录")}
            if cur_part is not None:
                cur_part["chapters"].append(len(chapters))
            chapters.append(cur_ch)
            cur_task = None
            # ★ 不 continue：本行 D 列可能已有知识点

        # ── 任务 ──
        m2 = re.match(r"^(任务[一二三四五六七八九十]+)[:：](.+)$", C)
        if m2:
            tname = m2.group(2).strip()
            if cur_task is None or cur_task["name"] != tname:
                if cur_ch is not None:
                    exist = next((t for t in cur_ch["tasks"] if t["name"] == tname), None)
                    if exist is None:
                        exist = {"seq": cn2int(m2.group(1)[2:]), "name": tname,
                                 "knowledge": []}
                        cur_ch["tasks"].append(exist)
                    cur_task = exist
            # ★ 不 continue：本行 D 列就是知识点

        # ── 知识点 ──
        m3 = re.match(r"^(\d+)\.\s*(.+)$", D)
        if m3 and cur_ch is not None:
            kname = m3.group(2).strip()
            # 无任务层级的章（原附录）：每个知识点自成一个任务
            # ★ 知识点名变化时重置 cur_task，否则第二个知识点会被并进第一个任务
            if cur_ch["no_task_layer"]:
                if cur_task is None or cur_task["name"] != kname:
                    exist = next((t for t in cur_ch["tasks"] if t["name"] == kname), None)
                    if exist is None:
                        exist = {"seq": len(cur_ch["tasks"]) + 1, "name": kname,
                                 "knowledge": [], "flat": True}
                        cur_ch["tasks"].append(exist)
                    cur_task = exist
            if cur_task is None:
                continue
            # 同名知识点多行 -> 合并核心内容
            for k in cur_task["knowledge"]:
                if k["name"] == kname:
                    if E:
                        k["entries"].append(E)
                    break
            else:
                cur_task["knowledge"].append({
                    "seq": len(cur_task["knowledge"]) + 1, "name": kname,
                    "entries": [E] if E else []})

    # ── 生成编码 ──
    for ci, ch in enumerate(chapters, 1):
        ch["code_suffix"] = f"C{ci:02d}"
        for ti, t in enumerate(ch["tasks"], 1):
            t["code_suffix"] = f"{ch['code_suffix']}-T{ti:02d}"
            for ki, k in enumerate(t["knowledge"], 1):
                k["code"] = f"{COURSE_CODE}-{t['code_suffix']}-K{ki:02d}"
    return parts, chapters


def main() -> int:
    parts, chapters = parse()

    n_ch = len(chapters)
    n_task = sum(len(c["tasks"]) for c in chapters)
    n_kp = sum(len(t["knowledge"]) for c in chapters for t in c["tasks"])
    n_entry = sum(len(k["entries"]) for c in chapters
                  for t in c["tasks"] for k in t["knowledge"])
    n_code = n_ch + n_task + n_kp

    print("=" * 76)
    print(f"知识点编码清单 · {COURSE_NAME}（{COURSE_CODE}）")
    print("=" * 76)
    for i, p in enumerate(parts, 1):
        print(f"  第{i}篇 {p['name']}（{len(p['chapters'])} 章）")
    print(f"\n篇 {len(parts)} / 章 {n_ch} / 任务 {n_task} / 知识点 {n_kp}")
    print(f"编码总数 {n_code}（{n_ch} 章 + {n_task} 任务 + {n_kp} 知识点）")
    print(f"知识库条目 {n_entry}")
    print()

    OUT_DIR.mkdir(parents=True, exist_ok=True)

    # ── Markdown ──
    md = [f"# {COURSE_NAME} · 知识点编码清单\n",
          f"课程编码：`{COURSE_CODE}`　编码格式：`{COURSE_CODE}-C01-T01-K01`\n",
          f"共 {len(parts)} 篇、{n_ch} 章、{n_task} 个任务、{n_kp} 个知识点"
          f"（编码 {n_code} 条，知识库条目 {n_entry} 条）\n",
          "\n编码规则：`课程码-C两位章号-T两位任务号-K两位知识点序号`。"
          "补零是为了字符串排序正确（不补零时 `C10` 会排在 `C2` 前面）。\n"]
    for p in parts:
        md.append(f"\n## {p['raw']}\n")
        for chi in p["chapters"]:
            ch = chapters[chi]
            md.append(f"\n### `{ch['code_suffix']}`　第{ch['seq']}章 {ch['name']}\n")
            md.append("\n| 编码 | 任务 | 知识点 | 知识库条目 |")
            md.append("\n| --- | --- | --- | --- |")
            for t in ch["tasks"]:
                for k in t["knowledge"]:
                    md.append(f"\n| `{k['code']}` | {t['name']} | {k['name']} "
                              f"| {len(k['entries'])} |")
            md.append("")
    (OUT_DIR / "知识点编码清单.md").write_text("\n".join(md), encoding="utf-8")

    # ── CSV ──
    with (OUT_DIR / "知识点编码清单.csv").open("w", encoding="utf-8-sig",
                                               newline="") as f:
        w = csv.writer(f)
        w.writerow(["课程编码", "课程名", "篇", "章编码", "章号", "章名",
                    "任务编码", "任务号", "任务名", "知识点编码", "知识点序号",
                    "知识点名", "知识库条目数"])
        for ch in chapters:
            part_raw = ch["part"]["raw"] if ch["part"] else ""
            for t in ch["tasks"]:
                for k in t["knowledge"]:
                    w.writerow([COURSE_CODE, COURSE_NAME, part_raw,
                                ch["code_suffix"], ch["seq"], ch["name"],
                                t["code_suffix"], t["seq"], t["name"],
                                k["code"], k["seq"], k["name"], len(k["entries"])])

    # ── JSON ──
    tree = {
        "course": {"code": COURSE_CODE, "name": COURSE_NAME,
                   "grade_band": GRADE_BAND},
        "parts": [{"name": p["raw"],
                   "chapters": [chapters[i]["code_suffix"] for i in p["chapters"]]}
                  for p in parts],
        "chapters": [
            {"code": ch["code_suffix"], "seq": ch["seq"], "name": ch["name"],
             "part": ch["part"]["raw"] if ch["part"] else None,
             "no_task_layer": ch["no_task_layer"],
             "tasks": [
                 {"code": t["code_suffix"], "seq": t["seq"], "name": t["name"],
                  "knowledge": [
                      {"code": k["code"], "seq": k["seq"], "name": k["name"],
                       "entries": k["entries"]}
                      for k in t["knowledge"]]}
                 for t in ch["tasks"]]}
            for ch in chapters],
    }
    (OUT_DIR / "知识点树.json").write_text(
        json.dumps(tree, ensure_ascii=False, indent=2), encoding="utf-8")

    print("生成:")
    for fn in ("知识点编码清单.md", "知识点编码清单.csv", "知识点树.json"):
        pp = OUT_DIR / fn
        print(f"  {fn:<22} {pp.stat().st_size/1024:>7.1f} KB")

    print()
    print("=" * 76)
    print("各章编码与知识点数")
    print("=" * 76)
    for ch in chapters:
        kp = sum(len(t["knowledge"]) for t in ch["tasks"])
        tag = "（无任务层级）" if ch["no_task_layer"] else ""
        print(f"  {ch['code_suffix']}  第{ch['seq']}章 {ch['name']:<12} "
              f"任务 {len(ch['tasks']):>2}  知识点 {kp:>2} {tag}")
    print()
    print("=" * 76)
    print("第六章编码明细（原附录）")
    print("=" * 76)
    for ch in chapters:
        if ch["no_task_layer"]:
            for t in ch["tasks"]:
                for k in t["knowledge"]:
                    print(f"  {k['code']:<24} {k['name']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
