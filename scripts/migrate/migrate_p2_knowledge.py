"""知识体系包生成（migrate_p2）。

输入：`docs/知识点清单.xlsx`
输出：`docs/教材/知识体系包.json`

★ 结构（4 层 + 知识库初稿）：

    篇（仅作为 chapter 的属性，不独立成层）
     └─ 章 chapter
         └─ 任务 task          ← 组卷的最小粒度
             └─ 知识点 knowledge_node
                 └─ 核心内容 → 知识库条目 knowledge_entry

★ 解析要点（都是从实测数据里踩出来的）：
  1. 表头在第 1 行，列位是 C(索引2)=章节、D(索引3)=知识点、E(索引4)=核心内容。
  2. **每行同时含「上级名称 + 知识点」**，所以上级分支不能 `continue`，
     否则知识点全丢。
  3. 「附录：润滑与密封」算正文 → 编为第六章，归第三篇。
     该章无任务层级，每个知识点自成一个任务。

用法：
    python scripts/migrate/migrate_p2_knowledge.py [--dry-run]
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from common import (  # noqa: E402
    COURSE_CODE,
    COURSE_NAME,
    GRADE_BAND,
    PACK_KNOWLEDGE,
    SRC_KNOWLEDGE,
    banner,
    file_sha256,
    now_iso,
    write_json,
)

from openpyxl import load_workbook  # noqa: E402

CN_NUM = "一二三四五六七八九十"

#: 列索引（0 基）
COL_CHAPTER = 2      # C 列：章节 / 篇 / 附录标记
COL_KNOWLEDGE = 3    # D 列：知识点
COL_CONTENT = 4      # E 列：核心内容


def cn2int(s: str) -> int:
    """中文数字 → 阿拉伯数字（支持一 ~ 二十）。"""
    if s == "十":
        return 10
    if s.startswith("十"):
        return 10 + CN_NUM.index(s[1]) + 1
    if "十" in s:
        a, _, b = s.partition("十")
        return (CN_NUM.index(a) + 1) * 10 + (CN_NUM.index(b) + 1 if b else 0)
    return CN_NUM.index(s) + 1


def s(v) -> str:
    return str(v).strip() if v is not None else ""


def parse() -> dict:
    rows = list(load_workbook(SRC_KNOWLEDGE, data_only=True)["Sheet1"]
                .iter_rows(values_only=True))

    parts: list[dict] = []
    chapters: list[dict] = []
    cur_part = cur_ch = cur_task = None
    no_task_layer = False

    for r in rows[1:]:
        C, D, E = s(r[COL_CHAPTER]), s(r[COL_KNOWLEDGE]), s(r[COL_CONTENT])

        # ── 篇（仅属性）──
        if re.match(r"^第.+篇[:：]", C):
            cur_part = {"raw": C, "name": C.split("：", 1)[1].strip(), "chapters": []}
            parts.append(cur_part)
            cur_ch = cur_task = None
            continue

        # ── 章（含原「附录」，按确认并入章序列）──
        if C.startswith("附录"):
            ch_name = C.split("：", 1)[1].strip() if "：" in C else "附录"
            is_appendix = True
        else:
            m = re.match(r"^(第[一二三四五六七八九十]+章)[:：](.+)$", C)
            ch_name = m.group(2).strip() if m else None
            is_appendix = False

        if ch_name and not any(c["name"] == ch_name for c in chapters):
            cur_ch = {
                "seq": len(chapters) + 1,
                "name": ch_name,
                "code": f"{COURSE_CODE}-C{len(chapters) + 1:02d}",
                "part": cur_part["raw"] if cur_part else None,
                "part_seq": parts.index(cur_part) + 1 if cur_part else None,
                "part_name": cur_part["name"] if cur_part else "",
                "no_task_layer": is_appendix,
                "tasks": [],
            }
            if cur_part is not None:
                cur_part["chapters"].append(len(chapters))
            chapters.append(cur_ch)
            cur_task = None
            # ★ 不 continue：本行 D 列可能已有知识点

        # ── 任务 ──
        m2 = re.match(r"^(任务[一二三四五六七八九十]+)[:：](.+)$", C)
        if m2:
            tname = m2.group(2).strip()
            if cur_task is None or cur_task["name"] != tname:
                if cur_ch is not None:
                    exist = next(
                        (t for t in cur_ch["tasks"] if t["name"] == tname), None
                    )
                    if exist is None:
                        tseq = cn2int(m2.group(1)[2:])
                        exist = {
                            "seq": tseq,
                            "name": tname,
                            "code": f"{cur_ch['code']}-T{tseq:02d}",
                            "knowledge": [],
                        }
                        cur_ch["tasks"].append(exist)
                    cur_task = exist
            # ★ 不 continue：本行 D 列就是知识点

        # ── 知识点 ──
        m3 = re.match(r"^(\d+)\.\s*(.+)$", D)
        if m3 and cur_ch is not None:
            kname = m3.group(2).strip()
            # 无任务层级的章（原附录）：每个知识点自成一个任务
            if cur_ch["no_task_layer"]:
                if cur_task is None or cur_task["name"] != kname:
                    exist = next(
                        (t for t in cur_ch["tasks"] if t["name"] == kname), None
                    )
                    if exist is None:
                        tseq = len(cur_ch["tasks"]) + 1
                        exist = {
                            "seq": tseq,
                            "name": kname,
                            "code": f"{cur_ch['code']}-T{tseq:02d}",
                            "knowledge": [],
                            "flat": True,
                        }
                        cur_ch["tasks"].append(exist)
                    cur_task = exist
            if cur_task is None:
                continue
            # 同名知识点多行 → 合并核心内容
            for k in cur_task["knowledge"]:
                if k["name"] == kname:
                    if E:
                        k["entries"].append(E)
                    break
            else:
                kseq = len(cur_task["knowledge"]) + 1
                cur_task["knowledge"].append({
                    "seq": kseq,
                    "name": kname,
                    "code": f"{cur_task['code']}-K{kseq:02d}",
                    "entries": [E] if E else [],
                })

    n_task = sum(len(c["tasks"]) for c in chapters)
    n_kp = sum(len(t["knowledge"]) for c in chapters for t in c["tasks"])
    n_entry = sum(len(k["entries"]) for c in chapters
                  for t in c["tasks"] for k in t["knowledge"])

    return {
        "meta": {
            "source_file": SRC_KNOWLEDGE.name,
            "source_sha256": file_sha256(SRC_KNOWLEDGE),
            "exported_at": now_iso(),
            "course_code": COURSE_CODE,
            "counts": {
                "parts": len(parts),
                "chapters": len(chapters),
                "tasks": n_task,
                "knowledge_nodes": n_kp,
                "knowledge_entries": n_entry,
            },
        },
        "course": {
            "code": COURSE_CODE,
            "name": COURSE_NAME,
            "grade_band": GRADE_BAND,
            "level": "SECONDARY",
            "style_config": {
                "context_required": True,
                "context_examples": [
                    "宝武集团", "比亚迪", "减速器维护", "数控机床", "东方电气",
                ],
                "context_rule": (
                    "命题必须嵌入真实岗位/企业情境，不得使用「某物体」「小明」"
                    "这类去情境化表述"
                ),
                "terminology_strict": True,
                "allow_drawing_question": False,
                "question_style": "TASK_DRIVEN",
            },
        },
        "parts": [{"seq": i + 1, "name": p["name"], "raw": p["raw"]}
                  for i, p in enumerate(parts)],
        "chapters": chapters,
    }


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true", help="只解析不写文件")
    args = ap.parse_args()

    banner("migrate_p2 · 知识体系包生成")
    pack = parse()
    c = pack["meta"]["counts"]

    print(f"源文件   {SRC_KNOWLEDGE.name}")
    print(f"篇       {c['parts']}")
    print(f"章       {c['chapters']}")
    print(f"任务     {c['tasks']}")
    print(f"知识点   {c['knowledge_nodes']}")
    print(f"知识库   {c['knowledge_entries']} 条核心内容")
    print()
    print("各章明细：")
    for ch in pack["chapters"]:
        nkp = sum(len(t["knowledge"]) for t in ch["tasks"])
        tag = "（无任务层级）" if ch["no_task_layer"] else ""
        print(f"  {ch['code']}  第{ch['seq']}章 {ch['name']:<12} "
              f"任务 {len(ch['tasks']):>2}  知识点 {nkp:>2} {tag}")

    # ── 完整性自检 ──
    problems: list[str] = []
    if c["knowledge_nodes"] == 0:
        problems.append("知识点为 0（解析失败）")
    if c["tasks"] == 0:
        problems.append("任务为 0（解析失败）")
    codes = [k["code"] for ch in pack["chapters"]
             for t in ch["tasks"] for k in t["knowledge"]]
    if len(codes) != len(set(codes)):
        problems.append(f"知识点编码重复：{len(codes)} vs {len(set(codes))}")

    print()
    if problems:
        print("自检失败：")
        for p in problems:
            print(f"  ✗ {p}")
        return 2
    print("自检通过 ✓")

    if args.dry_run:
        print("\n（--dry-run，未写入文件）")
    else:
        write_json(PACK_KNOWLEDGE, pack)
        size = PACK_KNOWLEDGE.stat().st_size / 1024
        print(f"\n已写入 {PACK_KNOWLEDGE}  ({size:.1f} KB)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
