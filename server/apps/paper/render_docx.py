"""试卷 Word 渲染（`P1-60` / 契约 v3，**可编辑**）。

★★ 几何来自 `layout.py`；Word 是流式文档，做**等价映射**：

    · 卷头标题/注意事项 = 段落；考生信息 = 带下划线制表位的段落
    · 题干 = 段落（已按 JSON 折行）
    · 选项 = 无边框表格（列数 = 该题 `option_cols`，1/2/4 列）
    · 配图 = inline picture
    · 空白区 = 固定行高空段
    · 页码 = 页脚 `fldSimple` 的 `PAGE` / `NUMPAGES`

  ★ 不用文本框、不用整页贴图；所有文字都能直接改。

★★ 中文字体必须同时设 `w:eastAsia`，否则换机器变方框。
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from typing import Any

from docx import Document
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH,
    WD_TAB_ALIGNMENT,
    WD_TAB_LEADER,
)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

from apps.paper import layout as L

BODY_FONT = "微软雅黑"
TITLE_FONT = "宋体"

_PT_TO_MM = 25.4 / 72.0


@dataclass
class DocxResult:
    data: bytes
    issues: list[dict[str, Any]] = field(default_factory=list)
    page_count: int = 0
    seqs_drawn: list[int] = field(default_factory=list)


def _set_run_font(run, *, name: str, size_pt: float, bold: bool = False):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:cs"), name)


def _style_paragraph(p, *, cfg: dict, space_after_pt: float = 0.0,
                     line_spacing: float | None = None, align=None):
    """统一段落间距：★ v3.1 段前/段后一律 0（行距仍按配置）。"""
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after_pt)
    pf.line_spacing = float(
        line_spacing if line_spacing is not None
        else cfg.get("line_spacing", 1.5))
    if align is not None:
        p.alignment = align


def _add_paragraph(doc, text: str, *, size: float, cfg: dict, bold=False,
                   align=None, space_after=0.0, font=BODY_FONT):
    p = doc.add_paragraph()
    _style_paragraph(p, cfg=cfg, space_after_pt=space_after, align=align)
    r = p.add_run(text)
    _set_run_font(r, name=font, size_pt=size, bold=bold)
    return p


def _add_student_lines(doc, header: dict, *, cfg: dict) -> None:
    """考生信息：标签 + **带下划线制表位**（不是 `___` 字符）。"""
    items = list(header.get("student_items") or [])
    if not items:
        return
    size = float(header.get("student_items", [{}])[0].get(
        "font_pt", cfg["base_font_pt"]))
    by_line: dict[int, list[dict]] = {}
    for item in items:
        by_line.setdefault(int(item.get("line_no") or 1), []).append(item)
    for line_no in sorted(by_line):
        line_items = sorted(by_line[line_no],
                            key=lambda x: float(x.get("x_mm") or 0))
        p = doc.add_paragraph()
        _style_paragraph(p, cfg=cfg)
        for idx, item in enumerate(line_items):
            label = str(item.get("label") or "")
            r = p.add_run(label)
            _set_run_font(r, name=BODY_FONT, size_pt=size)
            tab = p.add_run()
            tab.add_tab()
            end = float(item.get("underline_x_mm") or 0) \
                + float(item.get("underline_w_mm") or 0)
            p.paragraph_format.tab_stops.add_tab_stop(
                Mm(end), WD_TAB_ALIGNMENT.LEFT, WD_TAB_LEADER.LINES)
            if idx < len(line_items) - 1:
                gap = p.add_run("　")
                _set_run_font(gap, name=BODY_FONT, size_pt=size)


def _add_field(paragraph, instr: str, *, size: float, font: str = BODY_FONT):
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instr)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:eastAsia"), font)
    rfonts.set(qn("w:ascii"), font)
    rfonts.set(qn("w:hAnsi"), font)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rpr.append(rfonts)
    rpr.append(sz)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    paragraph._p.append(fld)


def _add_footer_page_number(doc, *, cfg: dict) -> None:
    sec = doc.sections[0]
    p = sec.footer.paragraphs[0]
    p.alignment = (WD_ALIGN_PARAGRAPH.CENTER
                   if cfg.get("page_number_align") == "center"
                   else WD_ALIGN_PARAGRAPH.RIGHT)
    _style_paragraph(p, cfg=cfg)
    size = 8.0
    r = p.add_run("第 ")
    _set_run_font(r, name=BODY_FONT, size_pt=size)
    _add_field(p, " PAGE ", size=size)
    r = p.add_run(" 页 / 共 ")
    _set_run_font(r, name=BODY_FONT, size_pt=size)
    _add_field(p, " NUMPAGES ", size=size)
    r = p.add_run(" 页")
    _set_run_font(r, name=BODY_FONT, size_pt=size)


def _enable_update_fields(doc) -> None:
    try:
        settings = doc.settings.element
        el = OxmlElement("w:updateFields")
        el.set(qn("w:val"), "true")
        settings.append(el)
    except Exception:  # noqa: BLE001
        pass


def _as_png(data: bytes) -> bytes:
    if data[:8] == b"\x89PNG\r\n\x1a\n" or data[:2] == b"\xff\xd8":
        return data
    from PIL import Image as PILImage

    im = PILImage.open(io.BytesIO(data))
    if im.mode not in ("RGB", "RGBA"):
        im = im.convert("RGB")
    buf = io.BytesIO()
    im.save(buf, format="PNG")
    return buf.getvalue()


def render_paper_docx(
    *,
    title: str,
    pages: list[Any],
    items: dict[int, dict[str, Any]],
    config: dict[str, Any],
    page_width_mm: float,
    page_height_mm: float,
    margin_mm: float,
    usable_width_mm: float,
    header: dict[str, Any] | None = None,
    assets: dict[int, bytes] | None = None,
    subtitle: str = "",
) -> DocxResult:
    """把试卷渲染成可编辑的 `.docx`。"""
    cfg = config
    doc = Document()
    issues: list[dict[str, Any]] = []
    seqs: list[int] = []
    assets = assets or {}
    header = header or L.measure_header(
        title=title, subtitle=subtitle, config=cfg,
        usable_width_mm=usable_width_mm)

    sec = doc.sections[0]
    sec.page_width = Mm(page_width_mm)
    sec.page_height = Mm(page_height_mm)
    sec.left_margin = Mm(margin_mm)
    sec.right_margin = Mm(margin_mm)
    sec.top_margin = Mm(margin_mm)
    sec.bottom_margin = Mm(margin_mm)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(cfg["base_font_pt"])
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    # ★ 兜住未显式设置的空段 / 表格单元格
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    # ── 卷头 ──
    title_h = header.get("title") or {}
    for s in (title_h.get("lines") or []):
        _add_paragraph(doc, s, size=float(title_h.get("font_pt") or 16),
                       cfg=cfg, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                       font=TITLE_FONT)
    sub = header.get("subtitle") or {}
    if sub.get("text"):
        _add_paragraph(doc, sub["text"],
                       size=float(sub.get("font_pt") or 10.5), cfg=cfg,
                       align=WD_ALIGN_PARAGRAPH.CENTER)
    _add_student_lines(doc, header, cfg=cfg)
    extra = header.get("student_extra") or {}
    if extra.get("text"):
        _add_paragraph(doc, extra["text"],
                       size=float(extra.get("font_pt") or 10.5),
                       cfg=cfg)
    notes = header.get("notes") or {}
    for s in (notes.get("lines") or []):
        _add_paragraph(doc, s, size=float(notes.get("font_pt") or 9.5),
                       cfg=cfg)

    # ── 正文（流式；不按 PDf 页强制分页 —— 可编辑优先）──
    for page in pages:
        for b in page.blocks:
            if b.kind == "stem":
                for ln in (getattr(b, "title_lines", None) or []):
                    if ln:
                        _add_paragraph(
                            doc, ln,
                            size=float(cfg["base_font_pt"]) + 1.5,
                            cfg=cfg, bold=True)
                continue

            if b.kind == "q_stems":
                for ln in (getattr(b, "stem_lines", None) or []):
                    if ln:
                        _add_paragraph(
                            doc, ln, size=float(cfg["base_font_pt"]),
                            cfg=cfg)
                if b.seq is not None and b.seq not in seqs:
                    seqs.append(b.seq)
                continue

            if b.kind == "choice_row":
                _add_options_table(doc, b, cfg=cfg, seqs_out=seqs)
                continue

            if b.kind == "subjective":
                label = getattr(b, "seq_label", "") or ""
                if label:
                    _add_paragraph(doc, label,
                                   size=float(cfg["base_font_pt"]),
                                   cfg=cfg, bold=True)
                for ln in (getattr(b, "stem_lines", None) or []):
                    if ln:
                        _add_paragraph(
                            doc, ln, size=float(cfg["base_font_pt"]),
                            cfg=cfg)
                if b.seq is not None and b.seq not in seqs:
                    seqs.append(b.seq)
                continue

            if b.kind == "figure":
                _add_figure(doc, b, assets=assets, cfg=cfg, issues=issues)
                continue

            if b.kind == "blank":
                line_pt = L.BLANK_LINE_GAP_MM / _PT_TO_MM
                n = max(int(round(float(b.height_mm)
                                  / L.BLANK_LINE_GAP_MM)), 1)
                for _ in range(n):
                    p = doc.add_paragraph()
                    _style_paragraph(p, cfg=cfg)
                    p.paragraph_format.line_spacing = Pt(line_pt)
                    r = p.add_run("")
                    _set_run_font(r, name=BODY_FONT,
                                  size_pt=float(cfg["base_font_pt"]))
                continue

    if cfg.get("show_page_number", True):
        _add_footer_page_number(doc, cfg=cfg)
        _enable_update_fields(doc)

    buf = io.BytesIO()
    doc.save(buf)
    return DocxResult(data=buf.getvalue(), issues=issues,
                      page_count=len(pages), seqs_drawn=sorted(seqs))


def _add_options_table(doc, block, *, cfg, seqs_out) -> None:
    """选项 = 无边框表格；行/列 = 该题 `option_cols`。"""
    cols = int(getattr(block, "option_cols", 0) or 1)
    rows = max(len(getattr(block, "option_row_heights", None) or []) or 1, 1)
    t = doc.add_table(rows=rows, cols=cols)
    t.autofit = False
    col_w = float(getattr(block, "option_w_mm", 0) or 0)
    if col_w > 0:
        for ci in range(cols):
            t.columns[ci].width = Mm(col_w)

    cells = {
        (int(c.get("row") or 0), int(c.get("col") or 0)): c
        for c in (getattr(block, "option_cells", None) or [])
    }
    font_pt = float(getattr(block, "option_font_pt", 0)
                    or (float(cfg["base_font_pt"]) - 1.5))
    for ri in range(rows):
        for ci in range(cols):
            cell = t.cell(ri, ci)
            cell.text = ""
            p = cell.paragraphs[0]
            _style_paragraph(p, cfg=cfg)
            info = cells.get((ri, ci))
            if not info:
                continue
            for li, ln in enumerate(info.get("lines") or []):
                if li:
                    br = p.add_run()
                    br.add_break()
                r = p.add_run(ln)
                _set_run_font(r, name=BODY_FONT, size_pt=font_pt)
    if block.seq is not None and block.seq not in seqs_out:
        seqs_out.append(block.seq)


def _add_figure(doc, b, *, assets, cfg, issues) -> None:
    aid = getattr(b, "asset_id", None)
    data = assets.get(int(aid)) if aid else None
    if getattr(b, "missing", False) or not data:
        _add_paragraph(doc, "［配图缺失］", size=float(cfg["base_font_pt"]),
                       cfg=cfg, align=WD_ALIGN_PARAGRAPH.CENTER)
        return
    try:
        p = doc.add_paragraph()
        _style_paragraph(p, cfg=cfg,
                         align=WD_ALIGN_PARAGRAPH.CENTER)
        run = p.add_run()
        run.add_picture(
            io.BytesIO(_as_png(data)),
            width=Mm(float(b.width_mm)),
            height=Mm(float(b.image_h_mm)))
    except Exception as exc:  # noqa: BLE001
        issues.append({
            "code": "PAPER.FIGURE_UNREADABLE",
            "message": f"第 {getattr(b, 'seq', '?')} 题的配图无法写入 Word："
                       f"{exc}",
            "blocking": False,
        })
        _add_paragraph(doc, "［配图无法显示］", size=float(cfg["base_font_pt"]),
                       cfg=cfg, align=WD_ALIGN_PARAGRAPH.CENTER)

    cap_font = max(float(cfg["base_font_pt"]) - 1.5, 6.0)
    for ln in (getattr(b, "caption_lines", None) or []):
        _add_paragraph(doc, ln, size=cap_font, cfg=cfg,
                       align=WD_ALIGN_PARAGRAPH.CENTER)
